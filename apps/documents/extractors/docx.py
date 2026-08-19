import logging

import docx
from docx.opc.exceptions import OpcError

from apps.documents.extractors.base import (
    BaseExtractor,
    CorruptedFileError,
    EmptyDocumentError,
    ExtractionError,
)

logger = logging.getLogger(__name__)


class DOCXExtractor(BaseExtractor):
    """
    Text extractor implementation for Microsoft Word (.docx) documents.
    Uses python-docx.
    """

    def extract_text(self, file_path: str) -> str:
        try:
            doc = docx.Document(file_path)
            paragraphs_text = []

            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    paragraphs_text.append(p.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            paragraphs_text.append(cell.text.strip())

            full_text = "\n\n".join(paragraphs_text).strip()
            if not full_text:
                raise EmptyDocumentError("DOCX document contains no extractable text.")

            return full_text

        except EmptyDocumentError:
            raise
        except OpcError as opc_err:
            logger.error("Corrupted DOCX zip package in %s: %s", file_path, opc_err)
            raise CorruptedFileError(f"DOCX file package is corrupted: {opc_err}")
        except Exception as exc:
            logger.error("Unexpected DOCX extraction error for %s: %s", file_path, exc)
            raise ExtractionError(f"Failed to extract DOCX text: {exc}")
